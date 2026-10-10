from __future__ import annotations

import csv
import difflib
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]

F129 = "129篇文献"
F_DTF = "企业数字化转型 100篇UTD24"
F_DCB = "数字环境下的消费者决策 100篇UTD24"
F_HUMAN = "AI人机协同相关主题文献1"
F_AIEXT = "AI相关学术论文文献汇总2"

TOPICS = [
    "人机协作与算法采纳",
    "AI战略、竞争优势与商业模式",
    "算法管理、劳动与就业",
    "创造力、创新与知识工作",
    "AI公平、伦理、透明与治理",
    "平台、生态系统与市场竞争",
    "营销、服务与消费者响应",
    "数字消费者决策与行为",
    "IT能力、数字化转型与企业绩效",
    "运营、供应链与高风险决策",
    "隐私、数据治理与安全",
    "智能体、多智能体与前沿技术",
]

S129_MAP = {
    "S2": "人机协作与算法采纳",
    "S1": "AI战略、竞争优势与商业模式",
    "S3": "算法管理、劳动与就业",
    "S4": "创造力、创新与知识工作",
    "S6": "AI公平、伦理、透明与治理",
    "S8": "平台、生态系统与市场竞争",
    "S7": "营销、服务与消费者响应",
    "S5": "运营、供应链与高风险决策",
}

DTF_MAP = {
    "01": "IT能力、数字化转型与企业绩效",
    "02": "AI战略、竞争优势与商业模式",
    "03": "创造力、创新与知识工作",
    "04": "平台、生态系统与市场竞争",
    "05": "IT能力、数字化转型与企业绩效",
    "06": "IT能力、数字化转型与企业绩效",
    "07": "算法管理、劳动与就业",
    "08": "AI战略、竞争优势与商业模式",
    "09": "运营、供应链与高风险决策",
    "10": "隐私、数据治理与安全",
}

HUMAN_FALLBACK = {
    "主题 1": "人机协作与算法采纳",
    "主题 2": "人机协作与算法采纳",
    "主题 3": "AI公平、伦理、透明与治理",
    "主题 4": "算法管理、劳动与就业",
    "主题 5": "平台、生态系统与市场竞争",
}

TOPIC_QUESTIONS = {
    "人机协作与算法采纳": [
        "生成式 AI 进入专业决策后，人类委托、复核与纠错行为如何随任务风险、解释方式和责任归属动态变化？",
        "算法建议的长期使用如何改变人的判断能力、注意分配和错误结构，并在什么条件下从增强转为依赖？",
    ],
    "AI战略、竞争优势与商业模式": [
        "AI 能力如何与数据网络效应、组织互补资源和治理机制共同形成可持续竞争壁垒？",
        "生成式 AI 对组织边界和商业模式的影响在什么条件下只是效率提升，何时会触发战略重构？",
    ],
    "算法管理、劳动与就业": [
        "算法管理下员工的遵从、规避与抵抗如何共同演化，并进一步影响绩效、福祉和职业发展？",
        "生成式 AI 对岗位任务重组、技能溢价与职业轨迹的短期调整和长期影响分别是什么？",
    ],
    "创造力、创新与知识工作": [
        "人机共创何时带来真正的创意增量，何时导致创意同质化、路径依赖或责任模糊？",
        "大语言模型参与知识生产后，知识贡献、审核机制与专业共同体规范如何变化？",
    ],
    "AI公平、伦理、透明与治理": [
        "解释、披露与人工复核如何影响不同能力群体的信任、遵从和公平感知？",
        "当公平约束与预测效用冲突时，组织如何设计可审计、可问责的人机决策流程？",
    ],
    "平台、生态系统与市场竞争": [
        "平台如何通过算法包络、数据接口和 Agent 基础设施改变进入壁垒、互补者创新与竞争结构？",
        "平台治理规则如何同时影响互补者创新、消费者福利、数据共享激励和市场集中？",
    ],
    "营销、服务与消费者响应": [
        "AI 客服、数字主播和生成内容披露如何改变信任、购买转化和长期客户关系？",
        "拟人化、情绪表达和人类参与披露在什么边界下提升服务体验，何时反而引发戒备？",
    ],
    "数字消费者决策与行为": [
        "推荐、排序、评价与隐私披露如何共同塑造消费者考虑集、选择质量和福利分配？",
        "LLM 辅助搜索和个性化推荐是否降低搜索成本，还是形成新的操纵、偏差与信息茧房？",
    ],
    "IT能力、数字化转型与企业绩效": [
        "IT 投资如何经过组织能力、流程互补和治理机制转化为可持续绩效，关键时滞在哪里？",
        "资源受限的中小企业或传统行业如何形成可复制、可测量且不加剧负担的数字化路径？",
    ],
    "运营、供应链与高风险决策": [
        "人机协同在供应链、医疗和高风险运营中如何同时平衡效率、韧性与安全？",
        "预测算法与人类判断的不同整合顺序和权限配置，如何影响库存、调度和临床决策？",
    ],
    "隐私、数据治理与安全": [
        "AI 的持续性、数据重用与跨场景溢出如何改变个人隐私边界和组织治理责任？",
        "隐私保护技术、解释机制与监管安排如何影响用户信任、企业采用和技术创新？",
    ],
    "智能体、多智能体与前沿技术": [
        "多智能体系统的协作、竞争、对齐与责任应由什么基础设施和治理结构承载？",
        "Agentic economy 中价值分配、平台权力与法律问责会如何演化，现有理论需要怎样扩展？",
    ],
}

TOPIC_METHODS = {
    "人机协作与算法采纳": "实验室/田野实验、企业操作日志、纵向追踪与行为计量。",
    "AI战略、竞争优势与商业模式": "企业面板、专利/招聘/投资数据、纵向案例、结构化访谈。",
    "算法管理、劳动与就业": "平台日志、HR 数据、自然实验、职业面板与田野访谈。",
    "创造力、创新与知识工作": "创意实验、文本嵌入、专利/R&D 数据和开源社区数据。",
    "AI公平、伦理、透明与治理": "审计实验、在线实验、组织案例、政策准实验与过程追踪。",
    "平台、生态系统与市场竞争": "平台数据、结构模型、博弈模型、监管变化准实验。",
    "营销、服务与消费者响应": "平台 A/B 测试、现场实验、服务交互数据与消费者实验。",
    "数字消费者决策与行为": "平台实验、结构模型、网站日志与消费者实验。",
    "IT能力、数字化转型与企业绩效": "企业面板、问卷、纵向案例与数字化项目数据。",
    "运营、供应链与高风险决策": "运营数据、田野实验、医院/制造案例与供应链建模。",
    "隐私、数据治理与安全": "隐私实验、政策准实验、平台数据与安全事件数据。",
    "智能体、多智能体与前沿技术": "模拟实验、博弈/机制设计、治理案例与平台实证数据。",
}


def clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("\xa0", " ")).strip()


def normalize_title(value: str) -> str:
    value = unicodedata.normalize("NFKC", value).lower()
    value = re.sub(r"[^\w\u4e00-\u9fff]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def parse_number(name: str) -> int | None:
    match = re.match(r"^(\d{1,3})", name)
    return int(match.group(1)) if match else None


def source_folder(relative_path: str) -> str:
    return relative_path.split("\\", 1)[0].split("/", 1)[0]


def extract_docx_index(path: Path, prefix: str) -> list[dict]:
    document = Document(path)
    paragraphs = [clean_text(paragraph.text) for paragraph in document.paragraphs]
    records: list[dict] = []
    current_theme = ""
    theme_re = re.compile(r"^(\d{2})\s+(.+)$")
    item_re = re.compile(rf"^({prefix}\d{{3}})\s+(.+)$")

    for index, text in enumerate(paragraphs):
        theme_match = theme_re.match(text)
        if (
            theme_match
            and 1 <= int(theme_match.group(1)) <= 10
            and "｜" not in text
            and not text.startswith("10大主题")
            and not theme_match.group(2).startswith("篇")
        ):
            current_theme = f"{theme_match.group(1)} {theme_match.group(2).strip()}"

        item_match = item_re.match(text)
        if not item_match:
            continue

        titles = [item_match.group(2).strip()]
        for next_text in paragraphs[index + 1 : index + 5]:
            if not next_text:
                continue
            if re.match(
                r"^(研究焦点|主要贡献|机制线索|边界条件|期刊|作者|附\s|"
                + prefix
                + r"\d{3})",
                next_text,
            ):
                break
            if re.search(r"[A-Za-z]", next_text) and len(next_text) <= 260:
                titles.append(next_text)
            else:
                break

        records.append(
            {
                "id": item_match.group(1),
                "titles": titles,
                "source_theme": current_theme,
            }
        )

    deduplicated: dict[str, dict] = {}
    for record in records:
        existing = deduplicated.get(record["id"])
        if existing is None or (
            not existing["source_theme"] and record["source_theme"]
        ):
            deduplicated[record["id"]] = record
    return list(deduplicated.values())


def extract_129_map(path: Path) -> dict[int, dict]:
    document = Document(path)
    for table in document.tables:
        if len(table.rows) < 2:
            continue
        headers = [clean_text(cell.text) for cell in table.rows[0].cells]
        if headers[:3] == ["编号", "中文标题", "英文标题"] and "主题" in headers:
            theme_index = headers.index("主题")
            result: dict[int, dict] = {}
            for row in table.rows[1:]:
                cells = [clean_text(cell.text) for cell in row.cells]
                try:
                    paper_id = int(cells[0])
                except (TypeError, ValueError):
                    continue
                result[paper_id] = {
                    "title": cells[1],
                    "source_theme": cells[theme_index],
                }
            return result
    raise RuntimeError("Could not find the 129-paper classification table.")


def extract_human_theme_map(path: Path) -> dict[int, str]:
    text = path.read_text(encoding="utf-8")
    matches = list(
        re.finditer(
            r"(?m)^### 主题\s+(\d+)：(.+?)（[^\n]*）\s*$",
            text,
        )
    )
    result: dict[int, str] = {}
    for index, match in enumerate(matches):
        theme_number = match.group(1)
        theme_name = match.group(2).strip()
        body_end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        body = text[match.end() : body_end]
        for paper_id_text in re.findall(
            r"(?<![\dA-Za-z])(\d{1,2})(?![\dA-Za-z])",
            body,
        ):
            paper_id = int(paper_id_text)
            if 1 <= paper_id <= 79:
                result.setdefault(paper_id, f"主题 {theme_number}：{theme_name}")
    return result


def classify_title(title: str) -> str:
    text = normalize_title(title)
    has = lambda *terms: any(term in text for term in terms)

    if has(
        "agent",
        "multi-agent",
        "multi ai",
        "llm",
        "large language model",
        "multimodal",
        "reinforcement learning",
        "lora",
        "tensorflow",
        "video understanding",
        "foundation model",
    ):
        return "智能体、多智能体与前沿技术"

    if has(
        "privacy",
        "data sharing",
        "data silo",
        "information security",
        "cybersecurity",
        "security",
        "surveillance",
        "risk information",
    ):
        return "隐私、数据治理与安全"

    if has(
        "supply chain",
        "inventory",
        "warehouse",
        "logistics",
        "healthcare",
        "diagnosis",
        "medical",
        "hospital",
        "traceability",
        "blockchain",
        "operational efficiency",
        "operations management",
    ):
        return "运营、供应链与高风险决策"

    if has(
        "employee",
        "worker",
        "workplace",
        "labor",
        "labour",
        "job",
        "hiring",
        "resume",
        "work stress",
        "wellbeing",
        "inclusion",
        "gig economy",
        "algorithmic management",
        "delegate",
        "delegation",
    ):
        return "算法管理、劳动与就业"

    if has(
        "customer",
        "consumer",
        "service",
        "chatbot",
        "streamer",
        "livestream",
        "live streaming",
        "advertising",
        "advertisement",
        "retail",
        "purchase",
        "brand",
        "luxury",
        "review",
        "word of mouth",
        "social media",
        "loyalty",
    ):
        return "营销、服务与消费者响应"

    if has(
        "search",
        "choice",
        "consideration",
        "personalization",
        "recommend",
        "impulse",
        "blind box",
        "subscription",
        "omnichannel",
        "e-commerce",
        "online shopping",
    ):
        return "数字消费者决策与行为"

    if has(
        "innovation",
        "creativity",
        "creative",
        "knowledge",
        "patent",
        "r&d",
        "idea generation",
    ):
        return "创造力、创新与知识工作"

    if has(
        "fairness",
        "bias",
        "ethics",
        "ethical",
        "explainable",
        "transparency",
        "accountability",
        "governance",
        "regulation",
        "regulating",
        "law",
        "alignment",
        "autonomous anomaly",
    ):
        return "AI公平、伦理、透明与治理"

    if has(
        "platform",
        "ecosystem",
        "market",
        "competition",
        "collusion",
        "pricing",
        "capitalism",
        "digital economy",
    ):
        return "平台、生态系统与市场竞争"

    if has(
        "digital transformation",
        "it capability",
        "firm performance",
        "enterprise system",
        "erp",
        "digital strategy",
        "organizational agility",
        "information technology",
        "technology mediated",
        "technology-mediated",
        "firm growth",
        "acquisition",
    ):
        return "IT能力、数字化转型与企业绩效"

    if has(
        "human-ai",
        "human ai",
        "human machine",
        "algorithm aversion",
        "algorithm reliance",
        "ai assistance",
        "decision support",
        "explainable artificial intelligence",
        "trust in artificial intelligence",
        "loss aversion",
        "automated vehicle",
        "autonomous driving",
        "self-driving",
        "drone",
    ):
        return "人机协作与算法采纳"

    if has(
        "strategy",
        "competitive advantage",
        "business model",
        "corporate",
        "entrepreneur",
        "finance",
        "financial",
        "stock price",
        "forecast",
    ):
        return "AI战略、竞争优势与商业模式"

    if has("social media", "enterprise social media", "information ecosystem"):
        return "平台、生态系统与市场竞争"

    return "智能体、多智能体与前沿技术"


def best_index_match(file_title: str, records: list[dict]) -> tuple[dict | None, float]:
    normalized_file = normalize_title(file_title)
    best_record = None
    best_score = 0.0
    for record in records:
        for candidate in record["titles"]:
            score = difflib.SequenceMatcher(
                None,
                normalized_file,
                normalize_title(candidate),
            ).ratio()
            if score > best_score:
                best_record = record
                best_score = score
    return best_record, best_score


def add_rows(
    rows: list[dict],
    folder: str,
    source_theme: str,
    source_id: str,
    title: str,
    unified_theme: str,
    path: Path,
) -> None:
    rows.append(
        {
            "统一主题": unified_theme,
            "来源资料夹": folder,
            "来源主题": source_theme or "标题关键词自动归类",
            "来源编号": source_id,
            "论文标题": title,
            "文件名": path.name,
            "相对路径": str(path.relative_to(ROOT)),
        }
    )


def build_rows() -> tuple[list[dict], dict[str, int]]:
    rows: list[dict] = []
    diagnostics: dict[str, int] = {}

    map129 = extract_129_map(ROOT / "AI与管理129篇顶刊文献系统综述报告.docx")
    for path in sorted((ROOT / F129).glob("*.pdf")):
        paper_id = parse_number(path.name)
        record = map129.get(paper_id or -1, {})
        source_theme = record.get("source_theme", "未匹配")
        unified = S129_MAP.get(source_theme, classify_title(path.stem))
        add_rows(
            rows,
            F129,
            source_theme,
            str(paper_id or ""),
            record.get("title") or path.stem,
            unified,
            path,
        )

    dtf_records = extract_docx_index(
        ROOT
        / F_DTF
        / "100篇UTD24文献知识库完整版（Word）企业数字化转型(1).docx",
        "DTF",
    )
    matched = 0
    for path in sorted((ROOT / F_DTF).glob("*.pdf")):
        record, score = best_index_match(path.stem, dtf_records)
        if record and score >= 0.52 and record["source_theme"][:2] in DTF_MAP:
            matched += 1
            source_theme = record["source_theme"]
            unified = DTF_MAP[source_theme[:2]]
            title = record["titles"][0]
            source_id = record["id"]
        else:
            source_theme = "标题关键词自动归类"
            unified = classify_title(path.stem)
            title = path.stem
            source_id = ""
        add_rows(rows, F_DTF, source_theme, source_id, title, unified, path)
    diagnostics["DTF 模糊匹配"] = matched

    dcb_records = extract_docx_index(
        ROOT
        / F_DCB
        / "100篇UTD24文献知识库完整版（Word）数字环境下的消费者决策(1).docx",
        "DCB",
    )
    matched = 0
    for path in sorted((ROOT / F_DCB).glob("*.pdf")):
        record, score = best_index_match(path.stem, dcb_records)
        if record and score >= 0.52:
            matched += 1
            source_theme = record["source_theme"]
            title = record["titles"][0]
            source_id = record["id"]
        else:
            source_theme = "标题关键词自动归类"
            title = path.stem
            source_id = ""
        add_rows(
            rows,
            F_DCB,
            source_theme,
            source_id,
            title,
            "数字消费者决策与行为",
            path,
        )
    diagnostics["DCB 模糊匹配"] = matched

    human_themes = extract_human_theme_map(
        ROOT / F_HUMAN / "人机协同文献深度分析报告.md"
    )
    matched = 0
    for path in sorted((ROOT / F_HUMAN).glob("*.pdf")):
        paper_id = parse_number(path.name)
        source_theme = human_themes.get(paper_id or -1, "")
        if source_theme:
            matched += 1
            theme_key = source_theme[:4]
            unified = HUMAN_FALLBACK.get(theme_key, classify_title(path.stem))
            if theme_key == "主题 5":
                unified = classify_title(path.stem)
        else:
            unified = classify_title(path.stem)
        add_rows(
            rows,
            F_HUMAN,
            source_theme,
            str(paper_id or ""),
            path.stem,
            unified,
            path,
        )
    diagnostics["人机协同显式主题匹配"] = matched

    for path in sorted((ROOT / F_AIEXT).glob("*.pdf")):
        add_rows(
            rows,
            F_AIEXT,
            "标题关键词自动归类",
            "",
            path.stem,
            classify_title(path.stem),
            path,
        )

    return rows, diagnostics


def write_csv(rows: list[dict]) -> None:
    fields = [
        "统一主题",
        "来源资料夹",
        "来源主题",
        "来源编号",
        "论文标题",
        "文件名",
        "相对路径",
    ]
    with (ROOT / "论文主题分类总表.csv").open(
        "w",
        encoding="utf-8-sig",
        newline="",
    ) as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def representative_rows(rows: list[dict], topic: str, limit: int = 6) -> list[dict]:
    folder_priority = {
        F129: 0,
        F_DTF: 1,
        F_DCB: 2,
        F_HUMAN: 3,
        F_AIEXT: 4,
    }
    candidates = [row for row in rows if row["统一主题"] == topic]
    return sorted(
        candidates,
        key=lambda row: (
            folder_priority.get(row["来源资料夹"], 9),
            row["来源编号"] or row["论文标题"],
        ),
    )[:limit]


def write_navigation(rows: list[dict], diagnostics: dict[str, int]) -> None:
    counts = Counter(row["统一主题"] for row in rows)
    lines = [
        "# 论文主题选题导航",
        "",
        "本索引把 493 篇 PDF 统一映射为 12 个总主题，同时保留各资料夹的原始主题。"
        "完整逐篇清单见 `论文主题分类总表.csv`。",
        "",
        "## 主题总览",
        "",
        "| 统一主题 | 论文数 | 主要来源 | 适合优先思考的选题方向 |",
        "| --- | ---: | --- | --- |",
    ]

    for topic in TOPICS:
        topic_rows = [row for row in rows if row["统一主题"] == topic]
        source_counts = Counter(row["来源资料夹"] for row in topic_rows)
        source_text = "；".join(
            f"{folder} {count} 篇" for folder, count in source_counts.most_common()
        )
        lines.append(
            f"| {topic} | {counts[topic]} | {source_text} | "
            f"{TOPIC_QUESTIONS[topic][0]} |"
        )

    lines.extend(
        [
            "",
            "## 分组阅读与选题线索",
            "",
        ]
    )

    for topic in TOPICS:
        topic_rows = [row for row in rows if row["统一主题"] == topic]
        source_counts = Counter(row["来源资料夹"] for row in topic_rows)
        source_themes = sorted(
            {
                row["来源主题"]
                for row in topic_rows
                if row["来源主题"]
                and row["来源主题"] not in {"标题关键词自动归类", "未匹配"}
            }
        )
        lines.extend(
            [
                f"### {topic}",
                "",
                f"- 论文规模：{len(topic_rows)} 篇",
                "- 资料来源："
                + "；".join(
                    f"{folder} {count} 篇"
                    for folder, count in source_counts.most_common()
                ),
                "- 原始主题："
                + ("；".join(source_themes) if source_themes else "主要由标题关键词归类"),
                f"- 推荐方法：{TOPIC_METHODS[topic]}",
                "- 候选选题：",
            ]
        )
        lines.extend(f"  - {question}" for question in TOPIC_QUESTIONS[topic])
        lines.append("- 代表文献：")
        for row in representative_rows(topic_rows, topic):
            lines.append(
                f"  - {row['论文标题']}（{row['来源资料夹']} / {row['来源主题']}）"
            )
        lines.append("")

    lines.extend(
        [
            "## 使用建议",
            "",
            "1. 先在总表中筛选一个统一主题，再按来源主题缩小到具体研究传统。",
            "2. 优先选择你已经能接触数据或实验场景的主题，不要只看文献规模。",
            "3. 同一主题内优先寻找“结论冲突、边界不清、纵向证据不足、生成式 AI 尚未覆盖”的缝隙。",
            "4. 跨主题组合通常更容易形成新问题，例如“人机协作 × 供应链韧性”“算法管理 × 员工福祉”“数字消费者 × 隐私治理”。",
            "",
            "## 生成诊断",
            "",
        ]
    )
    lines.extend(f"- {key}：{value} 篇" for key, value in diagnostics.items())
    lines.append("")
    (ROOT / "论文主题选题导航.md").write_text(
        "\n".join(lines),
        encoding="utf-8",
    )


def main() -> None:
    rows, diagnostics = build_rows()
    write_csv(rows)
    write_navigation(rows, diagnostics)
    print(f"Generated {len(rows)} rows.")
    for topic in TOPICS:
        print(f"{topic}: {sum(row['统一主题'] == topic for row in rows)}")
    for key, value in diagnostics.items():
        print(f"{key}: {value}")


if __name__ == "__main__":
    main()
